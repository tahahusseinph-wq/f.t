"""تصدير الجرد إلى Excel منسّق، واستيراد المنتجات من Excel، وتصدير أي جدول."""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Callable, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from ftapp.core.paths import logo_path, sub_dir
from ftapp.core.utils import money
from ftapp.models import Category, Product, Supplier, User
from ftapp.services import (audit, catalog_service, currency_service, inventory_service,
                            settings_service as settings)
from ftapp.services.errors import ValidationError

BRAND_BLUE = "1565C0"
LIGHT_BLUE = "E3F2FD"
LOW_FILL = PatternFill("solid", fgColor="FFE0B2")
OUT_FILL = PatternFill("solid", fgColor="FFCDD2")
HEADER_FILL = PatternFill("solid", fgColor=BRAND_BLUE)
TOTAL_FILL = PatternFill("solid", fgColor=LIGHT_BLUE)
THIN = Side(style="thin", color="B0BEC5")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FONT = "Cairo"

# الأعمدة المتاحة للتصدير: المفتاح ← (العنوان، دالة القيمة، نوع)
Getter = Callable[[Session, Product], Any]


def _export_columns(session: Session) -> dict[str, tuple[str, Getter, str]]:
    cols: dict[str, tuple[str, Getter, str]] = {
        "code": ("الكود", lambda s, p: p.code, "text"),
        "name": ("اسم المنتج", lambda s, p: p.name, "text"),
        "barcode": ("الباركود", lambda s, p: p.barcode or "", "text"),
        "category": ("القسم", lambda s, p: catalog_service.category_path(p.category), "text"),
        "brand": ("الماركة", lambda s, p: p.brand, "text"),
        "model": ("الموديل", lambda s, p: p.model, "text"),
        "unit": ("الوحدة", lambda s, p: p.unit, "text"),
        "quantity": ("الكمية", lambda s, p: p.quantity, "qty"),
        "min_stock": ("حد التنبيه", lambda s, p: catalog_service.min_stock_for(s, p), "qty"),
        "cost_price": ("سعر التكلفة", lambda s, p: p.cost_price, "money"),
        "sale_price": ("سعر البيع", lambda s, p: p.sale_price, "money"),
        "margin": ("نسبة الربح %", lambda s, p: catalog_service.effective_margin(s, p), "percent"),
        "cost_value": ("قيمة التكلفة", lambda s, p: money(p.quantity * p.cost_price), "money"),
        "sale_value": ("القيمة بسعر البيع", lambda s, p: money(p.quantity * p.sale_price), "money"),
        "supplier": ("المورد", lambda s, p: p.supplier.name if p.supplier else "", "text"),
        "location": ("الموقع", lambda s, p: p.location, "text"),
        "status": ("الحالة", lambda s, p: {"ok": "متوفر", "low": "منخفض", "out": "نافد"}[catalog_service.stock_status(s, p)], "text"),
        "details": ("التفاصيل", lambda s, p: p.details, "text"),
    }
    for fld in catalog_service.list_fields(session, all_fields=True):
        def getter(s: Session, p: Product, fid: int = fld.id, f=fld) -> Any:
            fv = next((v for v in p.field_values if v.field_id == fid), None)
            return catalog_service.display_field_value(f, fv.value) if fv else ""
        cols[f"cf:{fld.id}"] = (fld.name, getter, "text")
    return cols


DEFAULT_EXPORT_COLUMNS = ["code", "name", "barcode", "brand", "model", "unit", "quantity", "min_stock",
                          "cost_price", "sale_price", "margin", "cost_value", "sale_value", "status"]


def export_column_choices(session: Session) -> list[tuple[str, str]]:
    return [(k, v[0]) for k, v in _export_columns(session).items()]


def _safe_sheet_name(name: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "-", name)[:28] or "قسم"
    candidate, n = base, 2
    while candidate in used:
        candidate = f"{base[:25]} {n}"
        n += 1
    used.add(candidate)
    return candidate


def _style_header(ws: Worksheet, row: int, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = 28


def _write_title(ws: Worksheet, session: Session, title: str, ncols: int, with_logo: bool = True) -> int:
    """يكتب ترويسة الورقة (الشعار والاسم والتاريخ) ويعيد رقم صف بداية الجدول."""
    ws.sheet_view.rightToLeft = True
    company = settings.get(session, "company")
    last = get_column_letter(max(ncols, 4))
    ws.merge_cells(f"A1:{last}1")
    ws["A1"] = company.get("name", "")
    ws["A1"].font = Font(name=FONT, bold=True, size=16, color=BRAND_BLUE)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 34
    ws.merge_cells(f"A2:{last}2")
    ws["A2"] = f"{title} — {datetime.now():%Y-%m-%d %H:%M}"
    ws["A2"].font = Font(name=FONT, size=12, color="455A64")
    ws["A2"].alignment = Alignment(horizontal="center")
    if with_logo and logo_path().exists():
        try:
            img = XLImage(str(logo_path()))
            img.width = img.height = 60
            ws.add_image(img, "A1")
            ws.row_dimensions[1].height = 48
        except Exception:
            pass
    return 4


def _format_cell(cell, kind: str, symbol: str) -> None:
    cell.border = BORDER
    cell.font = Font(name=FONT, size=10)
    cell.alignment = Alignment(vertical="center", horizontal="right" if kind == "text" else "center")
    if kind == "money":
        cell.number_format = f'#,##0.00 "{symbol}"'
    elif kind == "qty":
        cell.number_format = "#,##0.###"
    elif kind == "percent":
        cell.number_format = '0.0"%"'


def _autosize(ws: Worksheet, start_row: int) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows(min_row=start_row):
        for cell in row:
            if cell.value is not None:
                widths[cell.column] = max(widths.get(cell.column, 0), len(str(cell.value)))
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(max(10, w + 4), 60)


def _write_products(ws: Worksheet, session: Session, products: list[Product], keys: list[str],
                    cols: dict[str, tuple[str, Getter, str]], start: int, symbol: str) -> dict[str, float]:
    for i, key in enumerate(keys, 1):
        ws.cell(row=start, column=i, value=cols[key][0])
    _style_header(ws, start, len(keys))
    totals = {"qty": 0.0, "cost": 0.0, "sale": 0.0}
    r = start
    for p in products:
        r += 1
        status = catalog_service.stock_status(session, p)
        for i, key in enumerate(keys, 1):
            _, getter, kind = cols[key]
            cell = ws.cell(row=r, column=i, value=getter(session, p))
            _format_cell(cell, kind, symbol)
            if status == "low":
                cell.fill = LOW_FILL
            elif status == "out":
                cell.fill = OUT_FILL
        totals["qty"] += p.quantity
        totals["cost"] += p.quantity * p.cost_price
        totals["sale"] += p.quantity * p.sale_price
    # صف الإجمالي
    r += 1
    ws.cell(row=r, column=1, value="الإجمالي")
    for key, tkey in (("quantity", "qty"), ("cost_value", "cost"), ("sale_value", "sale")):
        if key in keys:
            ws.cell(row=r, column=keys.index(key) + 1, value=round(totals[tkey], 2))
    for i, key in enumerate(keys, 1):
        cell = ws.cell(row=r, column=i)
        _format_cell(cell, cols[key][2] if cell.value is not None and i > 1 else "text", symbol)
        cell.fill = TOTAL_FILL
        cell.font = Font(name=FONT, bold=True)
    ws.freeze_panes = ws.cell(row=start + 1, column=1)
    ws.auto_filter.ref = f"A{start}:{get_column_letter(len(keys))}{r - 1}"
    _autosize(ws, start)
    return totals


@dataclass
class ExportResult:
    folder: Path
    excel: Path
    pdf: Path | None = None
    products: int = 0


def export_inventory(session: Session, actor: User | None, base_dir: Path | str | None = None,
                     columns: list[str] | None = None,
                     include_inactive: bool = False, sheet_per_category: bool = True) -> ExportResult:
    cols = _export_columns(session)
    keys = [k for k in (columns or DEFAULT_EXPORT_COLUMNS) if k in cols]
    if not keys:
        raise ValidationError("اختر عموداً واحداً على الأقل")
    today = date.today().isoformat()
    folder = Path(base_dir or sub_dir("exports")) / f"جرد_{today}"
    folder.mkdir(parents=True, exist_ok=True)
    symbol = currency_service.base(session).symbol

    products, _ = catalog_service.search_products(session, active=None if include_inactive else True, limit=None,
                                                  order="code")
    cats = {c.id: c for c in session.scalars(select(Category))}

    def root_of(p: Product) -> Category | None:
        c = cats.get(p.category_id) if p.category_id else None
        while c is not None and c.parent_id and c.parent_id in cats:
            c = cats[c.parent_id]
        return c

    groups: dict[int | None, list[Product]] = {}
    for p in products:
        r = root_of(p)
        groups.setdefault(r.id if r else None, []).append(p)

    wb = Workbook()
    summary = wb.active
    summary.title = "الملخص"
    used = {"الملخص"}
    start = _write_title(summary, session, "ملخص جرد المستودع", 6)
    headers = ["القسم", "عدد الأصناف", "الكمية", "قيمة التكلفة", "القيمة بسعر البيع", "أصناف منخفضة/نافدة"]
    for i, h in enumerate(headers, 1):
        summary.cell(row=start, column=i, value=h)
    _style_header(summary, start, len(headers))
    r = start
    grand = {"items": 0, "qty": 0.0, "cost": 0.0, "sale": 0.0, "low": 0}
    for cid, items in sorted(groups.items(), key=lambda kv: cats[kv[0]].name if kv[0] in cats else "ي"):
        name = cats[cid].name if cid in cats else "بدون قسم"
        if sheet_per_category:
            ws = wb.create_sheet(_safe_sheet_name(name, used))
            st = _write_title(ws, session, f"جرد قسم: {name}", len(keys))
            totals = _write_products(ws, session, items, keys, cols, st, symbol)
        else:
            totals = {"qty": sum(p.quantity for p in items),
                      "cost": sum(p.quantity * p.cost_price for p in items),
                      "sale": sum(p.quantity * p.sale_price for p in items)}
        low = sum(1 for p in items if catalog_service.stock_status(session, p) != "ok")
        r += 1
        values = [name, len(items), totals["qty"], round(totals["cost"], 2), round(totals["sale"], 2), low]
        kinds = ["text", "qty", "qty", "money", "money", "qty"]
        for i, (v, k) in enumerate(zip(values, kinds), 1):
            _format_cell(summary.cell(row=r, column=i, value=v), k, symbol)
        grand["items"] += len(items)
        grand["qty"] += totals["qty"]
        grand["cost"] += totals["cost"]
        grand["sale"] += totals["sale"]
        grand["low"] += low
    r += 1
    values = ["الإجمالي", grand["items"], grand["qty"], round(grand["cost"], 2), round(grand["sale"], 2), grand["low"]]
    for i, (v, k) in enumerate(zip(values, ["text", "qty", "qty", "money", "money", "qty"]), 1):
        cell = summary.cell(row=r, column=i, value=v)
        _format_cell(cell, k, symbol)
        cell.fill = TOTAL_FILL
        cell.font = Font(name=FONT, bold=True)
    _autosize(summary, start)

    if not sheet_per_category:
        ws = wb.create_sheet("كل المنتجات")
        st = _write_title(ws, session, "جرد كامل", len(keys))
        _write_products(ws, session, products, keys, cols, st, symbol)

    excel = folder / f"جرد_المستودع_{today}.xlsx"
    wb.save(excel)
    audit.log(session, actor, "export", "inventory", None, file=str(excel), products=len(products))
    return ExportResult(folder=folder, excel=excel, products=len(products))


def export_table(path: Path | str, title: str, headers: list[str], rows: Iterable[Iterable[Any]],
                 session: Session | None = None, kinds: list[str] | None = None) -> Path:
    """تصدير أي جدول من الواجهة إلى Excel بنفس التنسيق."""
    wb = Workbook()
    ws = wb.active
    ws.title = _safe_sheet_name(title, set())
    symbol = ""
    if session is not None:
        start = _write_title(ws, session, title, len(headers), with_logo=False)
        symbol = currency_service.base(session).symbol
    else:
        ws.sheet_view.rightToLeft = True
        start = 1
    for i, h in enumerate(headers, 1):
        ws.cell(row=start, column=i, value=h)
    _style_header(ws, start, len(headers))
    r = start
    for row in rows:
        r += 1
        for i, v in enumerate(row, 1):
            kind = (kinds[i - 1] if kinds and i - 1 < len(kinds) else
                    ("qty" if isinstance(v, (int, float)) and not isinstance(v, bool) else "text"))
            _format_cell(ws.cell(row=r, column=i, value=v), kind, symbol)
    ws.freeze_panes = ws.cell(row=start + 1, column=1)
    if r > start:
        ws.auto_filter.ref = f"A{start}:{get_column_letter(len(headers))}{r}"
    _autosize(ws, start)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# =====================================================================
# الاستيراد
# =====================================================================

IMPORT_TARGETS: dict[str, str] = {
    "name": "اسم المنتج",
    "code": "الكود",
    "barcode": "الباركود",
    "category": "القسم",
    "brand": "الماركة",
    "model": "الموديل",
    "unit": "الوحدة",
    "cost_price": "سعر التكلفة",
    "sale_price": "سعر البيع",
    "margin": "نسبة الربح",
    "quantity": "الكمية",
    "min_stock": "حد التنبيه",
    "supplier": "المورد",
    "location": "الموقع",
    "details": "التفاصيل",
}

_ALIASES = {
    "name": ["اسم", "المنتج", "الصنف", "name", "product", "item", "description"],
    "code": ["كود", "رمز", "code", "sku", "ref"],
    "barcode": ["باركود", "barcode", "ean", "upc"],
    "category": ["قسم", "فئة", "تصنيف", "category", "group"],
    "brand": ["ماركة", "شركة", "brand", "make"],
    "model": ["موديل", "طراز", "model"],
    "unit": ["وحدة", "unit"],
    "cost_price": ["تكلفة", "شراء", "cost", "purchase"],
    "sale_price": ["بيع", "سعر", "price", "sale"],
    "margin": ["ربح", "نسبة", "margin", "markup"],
    "quantity": ["كمية", "عدد", "رصيد", "qty", "quantity", "stock"],
    "min_stock": ["حد", "تنبيه", "min", "reorder"],
    "supplier": ["مورد", "supplier", "vendor"],
    "location": ["موقع", "رف", "location", "shelf"],
    "details": ["تفاصيل", "وصف", "ملاحظات", "details", "notes"],
}


def read_preview(path: Path | str, sheet: str | None = None, max_rows: int = 20) -> dict[str, Any]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    # أول صف فيه نصوص كثيرة يُعتبر صف العناوين
    header_idx = 0
    for i, row in enumerate(rows[:10]):
        if sum(1 for v in row if isinstance(v, str) and v.strip()) >= max(2, len([v for v in row if v]) // 2):
            header_idx = i
            break
    headers = [str(v).strip() if v is not None else f"عمود {i + 1}" for i, v in enumerate(rows[header_idx] if rows else [])]
    data = [list(r) for r in rows[header_idx + 1:] if any(v not in (None, "") for v in r)]
    return {"sheets": wb.sheetnames, "headers": headers, "rows": data[:max_rows], "total": len(data),
            "header_row": header_idx}


def auto_map(headers: list[str], custom_fields: list[tuple[int, str]] | None = None) -> dict[int, str]:
    """يقترح ربط أعمدة الملف بحقول المنتج: {رقم العمود: المفتاح}."""
    mapping: dict[int, str] = {}
    used: set[str] = set()
    for i, h in enumerate(headers):
        hl = h.lower()
        best, score = None, 0.0
        for key, aliases in _ALIASES.items():
            if key in used:
                continue
            for a in aliases + [IMPORT_TARGETS[key]]:
                s = 1.0 if a in hl else SequenceMatcher(None, a, hl).ratio()
                if s > score:
                    best, score = key, s
        for fid, fname in custom_fields or []:
            s = SequenceMatcher(None, fname.lower(), hl).ratio()
            if s > score and f"cf:{fid}" not in used:
                best, score = f"cf:{fid}", s
        if best and score >= 0.6:
            mapping[i] = best
            used.add(best)
    return mapping


@dataclass
class ImportResult:
    created: int = 0
    updated: int = 0
    skipped: int = 0
    errors: list[tuple[int, str]] = field(default_factory=list)


def _num(v: Any) -> float | None:
    if v in (None, ""):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    text = str(v).replace(",", "").replace("٫", ".").strip()
    text = text.translate(str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789"))
    text = re.sub(r"[^\d.\-]", "", text)
    return float(text) if text not in ("", "-", ".") else None


def import_products(session: Session, actor: User | None, path: Path | str, mapping: dict[int, str],
                    sheet: str | None = None, update_existing: bool = True, create_categories: bool = True,
                    dry_run: bool = False) -> ImportResult:
    if "name" not in mapping.values() and "code" not in mapping.values():
        raise ValidationError("يجب ربط عمود اسم المنتج أو الكود على الأقل")
    preview = read_preview(path, sheet, max_rows=10**7)
    result = ImportResult()
    cats = {c.name.strip(): c for c in session.scalars(select(Category))}
    suppliers = {s.name.strip(): s for s in session.scalars(select(Supplier))}
    wh = inventory_service.default_warehouse(session).id

    for idx, row in enumerate(preview["rows"]):
        excel_row = preview["header_row"] + idx + 2
        values = {key: (row[col] if col < len(row) else None) for col, key in mapping.items()}
        sp = session.begin_nested()
        try:
            code = str(values.get("code") or "").strip()
            existing = catalog_service.find_by_code(session, code) if code else None
            if existing is None and values.get("barcode"):
                existing = catalog_service.find_by_code(session, str(values["barcode"]).strip())
            if existing and not update_existing:
                result.skipped += 1
                sp.rollback()
                continue
            data = catalog_service.product_input_from(existing) if existing else catalog_service.ProductInput(name="")
            if values.get("name"):
                data.name = str(values["name"]).strip()
            if code:
                data.code = code
            for key in ("barcode", "brand", "model", "unit", "location", "details"):
                if values.get(key) not in (None, ""):
                    v = values[key]
                    if key == "barcode" and isinstance(v, float) and v.is_integer():
                        v = int(v)
                    setattr(data, key, str(v).strip())
            for key in ("cost_price", "sale_price"):
                n = _num(values.get(key))
                if n is not None:
                    setattr(data, key, n)
            if _num(values.get("margin")) is not None:
                data.margin = _num(values.get("margin"))
            elif values.get("sale_price") not in (None, "") and "margin" not in mapping.values():
                data.price_locked = True
            if _num(values.get("min_stock")) is not None:
                data.min_stock = _num(values.get("min_stock"))
            cat_name = str(values.get("category") or "").strip()
            if cat_name:
                cat = cats.get(cat_name)
                if cat is None and create_categories:
                    cat = catalog_service.save_category(session, actor, cat_name)
                    cats[cat_name] = cat
                data.category_id = cat.id if cat else data.category_id
            sup_name = str(values.get("supplier") or "").strip()
            if sup_name:
                sup = suppliers.get(sup_name) or catalog_service.save_supplier(session, sup_name)
                suppliers[sup_name] = sup
                data.supplier_id = sup.id
            for key, v in values.items():
                if key.startswith("cf:") and v not in (None, ""):
                    data.custom_values[int(key[3:])] = v.isoformat() if isinstance(v, (date, datetime)) else v
            quantity = _num(values.get("quantity"))
            if existing:
                catalog_service.update_product(session, actor, existing.id, data, source="import")
                if quantity is not None:
                    inventory_service.set_quantity(session, actor, existing.id, wh, quantity, "استيراد من Excel")
                result.updated += 1
            else:
                if not data.name:
                    raise ValidationError("اسم المنتج فارغ")
                data.initial_quantity = quantity or 0
                catalog_service.create_product(session, actor, data, source="import")
                result.created += 1
            sp.commit()
        except Exception as exc:  # نسجل الخطأ ونكمل بقية الصفوف
            sp.rollback()
            result.errors.append((excel_row, str(exc)))
    if dry_run:
        session.rollback()
    else:
        audit.log(session, actor, "import", "product", None, created=result.created, updated=result.updated,
                  errors=len(result.errors))
    return result


def import_template(path: Path | str) -> Path:
    headers = [IMPORT_TARGETS[k] for k in ("code", "name", "barcode", "category", "brand", "model", "unit",
                                           "cost_price", "margin", "sale_price", "quantity", "min_stock",
                                           "supplier", "location", "details")]
    sample = [["", "سماعة بلوتوث", "", "إلكترونيات", "Xiaomi", "Buds 4", "قطعة", 12, 30, "", 20, 5, "", "A-3", ""]]
    return export_table(path, "قالب الاستيراد", headers, sample)
