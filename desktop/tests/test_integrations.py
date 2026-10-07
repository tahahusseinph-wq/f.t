import json
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from PySide6.QtCore import QUrl
from PySide6.QtGui import QTextDocument

from ftapp.services import (backup_service, catalog_service, excel_service, gemini_service, pdf_service,
                            sales_service, update_service)
from ftapp.services.errors import ValidationError


def test_excel_export_folder(db, admin, product_factory, tmp_path):
    cat = catalog_service.save_category(db, admin, "أدوات/كهرباء")
    fld = catalog_service.save_field(db, admin, "بلد المنشأ")
    product_factory("مفك", qty=3, min_stock=5, category_id=cat.id, custom_values={fld.id: "ألمانيا"})
    product_factory("شريط", qty=0)
    product_factory("لمبة", qty=40)
    res = excel_service.export_inventory(db, admin, tmp_path, columns=excel_service.DEFAULT_EXPORT_COLUMNS + [f"cf:{fld.id}"])
    assert res.folder.name.startswith("جرد_") and res.excel.exists()
    wb = load_workbook(res.excel)
    assert wb.sheetnames[0] == "الملخص" and len(wb.sheetnames) == 3
    assert wb["الملخص"].sheet_view.rightToLeft
    sheet = wb[wb.sheetnames[1]]
    values = [c.value for row in sheet.iter_rows() for c in row]
    assert "ألمانيا" in values and "بلد المنشأ" in values


def test_excel_import_roundtrip(db, admin, product_factory, tmp_path):
    existing = product_factory("قديم", qty=1, code="OLD-1")
    path = excel_service.export_table(tmp_path / "in.xlsx", "استيراد",
                                      ["الكود", "اسم الصنف", "القسم", "سعر الشراء", "نسبة الربح", "الكمية", "باركود"],
                                      [["OLD-1", "قديم معدل", "مطبخ", 5, 20, 7, ""],
                                       ["", "جديد", "مطبخ", "١٠", 50, 4, 6281234567890],
                                       ["", "", "", "", "", "", ""],
                                       ["", "", "مطبخ", "x", "", 2, ""]])
    preview = excel_service.read_preview(path)
    mapping = excel_service.auto_map(preview["headers"])
    assert set(mapping.values()) >= {"code", "name", "category", "cost_price", "margin", "quantity", "barcode"}
    res = excel_service.import_products(db, admin, path, mapping)
    db.commit()
    assert res.created == 1 and res.updated == 1 and len(res.errors) == 1
    db.refresh(existing)
    assert existing.name == "قديم معدل" and existing.quantity == 7 and existing.sale_price == 6
    new = catalog_service.find_by_code(db, "6281234567890")
    assert new and new.cost_price == 10 and new.sale_price == 15 and new.category.name == "مطبخ"


def test_backup_restore_encrypted(db, admin, product_factory, tmp_path):
    product_factory("قبل النسخ")
    db.commit()
    path = backup_service.create_backup(tmp_path, password="secret-pw")
    assert backup_service.is_encrypted(path)
    product_factory("بعد النسخ")
    db.commit()
    db.close()
    with pytest.raises(ValidationError):
        backup_service.restore_backup(path, "wrong")
    backup_service.restore_backup(path, "secret-pw")
    from ftapp.core import db as dbm
    s = dbm.new_session()
    names = [p.name for p in catalog_service.search_products(s)[0]]
    assert names == ["قبل النسخ"]
    s.close()


def test_invoice_pdf_and_labels(db, admin, product_factory, tmp_path):
    p = product_factory("سماعة", qty=5)
    inv = sales_service.create_sale(db, admin, sales_service.SaleRequest([sales_service.CartLine(p.id, 1)]))
    db.commit()
    for paper in ("A4", "80mm"):
        out = pdf_service.invoice_pdf(db, inv, tmp_path / f"{paper}.pdf", paper)
        assert out.read_bytes()[:4] == b"%PDF"
    lbl = pdf_service.labels_pdf(db, [(p, 3)], tmp_path / "l.pdf", sheet="a4")
    assert lbl.stat().st_size > 1000


def test_invoice_custom_details_facebook_and_shamcash(db, admin, product_factory):
    from ftapp.services import settings_service

    settings_service.update(db, "company", phone="0991234567", invoice_details="سجل تجاري 123\nواتساب 0999",
                            facebook_url="ft.trading", shamcash_account="SC-778899")
    p = product_factory("شاحن", qty=5)
    inv = sales_service.create_sale(db, admin, sales_service.SaleRequest(
        [sales_service.CartLine(p.id, 1)], payment_method="shamcash", notes="رقم عملية شام كاش: 555"))
    db.commit()
    assert pdf_service.facebook_link("ft.trading") == "https://www.facebook.com/ft.trading"
    assert pdf_service.facebook_link("facebook.com/x") == "https://facebook.com/x"
    for paper in ("A4", "80mm"):
        doc = pdf_service.build_invoice_document(db, inv, paper)
        text = doc.toPlainText()
        assert "سجل تجاري 123" in text and "واتساب 0999" in text and "0991234567" in text
        assert "SC-778899" in text and "تابعونا على فيسبوك" in text and "555" in text
        assert not doc.resource(QTextDocument.ResourceType.ImageResource, QUrl("fb")).isNull()


def test_commercial_invoice_customer_details_and_design_settings(db, admin, product_factory, tmp_path):
    from ftapp.services import settings_service

    settings_service.update(db, "invoice", sale_title="فاتورة تجارية", accent_color="#2E7D32",
                            terms="الكفالة سنة واحدة\nلا يُرد المبيع", payment_info="حساب بنكي 4455",
                            show_code=False, show_signatures=True)
    p = product_factory("راوتر", qty=5, code="RT-77")
    inv = sales_service.create_sale(db, admin, sales_service.SaleRequest(
        [sales_service.CartLine(p.id, 2)], customer_name="شركة النور", customer_phone="0933111222",
        customer_address="حلب - الجميلية"))
    db.commit()
    assert inv.customer_address == "حلب - الجميلية" and inv.customer_id is None
    text = pdf_service.build_invoice_document(db, inv, "A4").toPlainText()
    for expected in ("فاتورة تجارية", "شركة النور", "0933111222", "حلب - الجميلية", "الكفالة سنة واحدة",
                     "حساب بنكي 4455", "فاتورة إلى", "توقيع المستلم"):
        assert expected in text
    assert p.code not in text  # عمود الكود مخفي من الإعدادات
    assert "حلب - الجميلية" in pdf_service.build_invoice_document(db, inv, "80mm").toPlainText()
    assert pdf_service.invoice_pdf(db, inv, tmp_path / "c.pdf", "A4").read_bytes()[:4] == b"%PDF"


class _FakeModels:
    def __init__(self, payload):
        self.payload = payload
        self.calls = 0

    def generate_content(self, model, contents, config):
        self.calls += 1
        return SimpleNamespace(parsed=None, text=json.dumps(self.payload, ensure_ascii=False))


def test_gemini_details_cached_and_field_match(db, admin, monkeypatch):
    payload = {"description": "سماعة لاسلكية", "specs": [{"name": "البطارية", "value": "30 ساعة"}],
               "custom_fields": [{"name": "بلد المنشأ", "value": "الصين"}], "confidence": "high"}
    fake = _FakeModels(payload)
    monkeypatch.setattr(gemini_service, "_client", lambda key=None: SimpleNamespace(models=fake))
    d1 = gemini_service.fetch_product_details(db, "Buds 4", "Xiaomi", field_names=["بلد المنشأ", "البطارية"])
    d2 = gemini_service.fetch_product_details(db, "Buds 4", "Xiaomi", field_names=["بلد المنشأ", "البطارية"])
    assert fake.calls == 1 and d2.description == "سماعة لاسلكية"
    f1 = catalog_service.save_field(db, admin, "بلد المنشأ")
    f2 = catalog_service.save_field(db, admin, "البطارية")
    assert gemini_service.match_custom_fields(d1, [f1, f2]) == {f1.id: "الصين", f2.id: "30 ساعة"}


def test_gemini_error_translation():
    err = gemini_service._translate_error(Exception("400 API key not valid. API_KEY_INVALID"))
    assert "غير صالح" in str(err)
    err = gemini_service._translate_error(Exception("429 RESOURCE_EXHAUSTED"))
    assert "تجاوز" in str(err)


def test_update_check(monkeypatch):
    def fake_get(url, **kw):
        return SimpleNamespace(raise_for_status=lambda: None,
                               json=lambda: {"tag_name": "v9.0.0", "html_url": "https://x", "body": "جديد",
                                             "assets": [{"name": "app.apk", "browser_download_url": "https://x/a.apk"}]})
    monkeypatch.setattr(update_service.httpx, "get", fake_get)
    info = update_service.check("https://api.github.com/repos/o/r/releases/latest")
    assert info.is_newer and info.apk_url.endswith(".apk")


def test_product_warranty_on_invoice_and_view(db, admin, product_factory):
    p = product_factory("شاشة", qty=5, warranty="سنتان")
    q = product_factory("كبل", qty=5)
    inv = sales_service.create_sale(db, admin, sales_service.SaleRequest(
        [sales_service.CartLine(p.id, 1), sales_service.CartLine(q.id, 1)]))
    db.commit()
    assert [i.warranty for i in inv.items] == ["سنتان", ""]
    text = pdf_service.build_invoice_document(db, inv, "A4").toPlainText()
    assert "الكفالة" in text and "سنتان" in text and "بدون" in text
    assert "كفالة سنتان" in pdf_service.build_invoice_document(db, inv, "80mm").toPlainText()
    view = catalog_service.product_view(db, q, privileged=False)
    assert {"key": "warranty", "label": "الكفالة", "value": "بدون كفالة", "type": "text"} in view["fields"]


def test_migration_renames_company_and_updates_default_terms(tmp_data):
    import os

    import sqlalchemy as sa
    from alembic import command
    from alembic.config import Config

    from ftapp.core import db as dbm
    from ftapp.core.paths import migrations_dir

    dbm.dispose()
    dbm.init_engine(os.path.join(tmp_data, "old.db"))
    cfg = Config()
    cfg.set_main_option("script_location", str(migrations_dir()).replace("%", "%%"))
    with dbm.engine().begin() as conn:
        cfg.attributes["connection"] = conn
        command.upgrade(cfg, "0003")
        conn.execute(sa.text("INSERT INTO settings (key, value) VALUES ('company', :c), ('invoice', :i)"),
                     {"c": json.dumps({"name": "مجموعة فاروق الطعمة التجارية", "phone": "099"}, ensure_ascii=False),
                      "i": json.dumps({"terms": "البضاعة المباعة لا تُرد ولا تُستبدل إلا بموجب هذه الفاتورة وخلال 7 أيام."},
                                      ensure_ascii=False)})
    dbm.run_migrations()
    s = dbm.new_session()
    try:
        from ftapp.services import settings_service
        company = settings_service.get(s, "company")
        assert company["name"] == "مجموعة الطعمة التجارية" and company["phone"] == "099"
        assert settings_service.get(s, "invoice")["terms"] == "البضاعة التي تُباع لا تُرد ولا تُستبدل أبداً."
    finally:
        s.close()
        dbm.dispose()
